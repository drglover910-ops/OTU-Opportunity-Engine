"""Editable-text navy/gold résumé layout based on the owner's visual reference."""
import html,io,re
from reportlab.pdfgen import canvas
from reportlab.lib.colors import HexColor,white
from reportlab.platypus import Paragraph
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.pagesizes import letter
NAVY=HexColor('#09263B');GOLD=HexColor('#B89243');PALE=HexColor('#FBF5E5')
def render(data):
 out=io.BytesIO();c=canvas.Canvas(out,pagesize=letter);W,H=letter;page=0;y=0
 body=ParagraphStyle('body',fontName='Times-Roman',fontSize=10,leading=12.5,textColor=NAVY)
 side=ParagraphStyle('side',parent=body,fontName='OTURegular',fontSize=8.4,leading=11,textColor=white)
 def para(text,x,y,width,style=body):
  p=Paragraph(html.escape(str(text)),style);_,h=p.wrap(width,700);p.drawOn(c,x,y-h);return y-h-6
 def heading(text,x,y,width,dark=False):
  c.setFillColor(GOLD if dark else NAVY);c.circle(x+7,y-7,7,fill=1,stroke=0)
  c.setFillColor(NAVY if dark else GOLD);c.setFont('OTUBold',7);c.drawCentredString(x+7,y-9,text[0]);c.setLineWidth(.7)
  c.setFillColor(white if dark else NAVY);c.setFont('Times-Bold',10);c.drawString(x+21,y-11,text)
  c.setStrokeColor(GOLD);c.line(x+21,y-18,x+width,y-18);return y-28
 def sidebar(title,texts,top):
  # Measure every paragraph so panels expand for actual source content.
  heights=[]
  for text in texts:
   p=Paragraph(html.escape(str(text)),side);heights.append(p.wrap(158,700)[1]+7)
  h=36+sum(heights)+8
  if top-h<42:raise ValueError('Résumé sidebar needs a shorter verified selection. No clipped document was saved.')
  c.setFillColor(NAVY);c.roundRect(410,top-h,180,h,11,fill=1,stroke=0)
  sy=heading(title,420,top-10,160,True)
  for text in texts:sy=para(text,421,sy,158,side)
  return top-h-8
 def newpage():
  nonlocal page,y
  if page:c.showPage()
  page+=1;c.setFillColor(PALE);c.roundRect(7,7,W-14,H-14,15,fill=1,stroke=0);c.setFillColor(white);c.roundRect(12,12,W-24,H-24,12,fill=1,stroke=0)
  c.setStrokeColor(GOLD);c.setLineWidth(.8);c.roundRect(8,8,W-16,H-16,15,fill=0,stroke=1)
  c.setFillColor(NAVY);path=c.beginPath();path.moveTo(424,H-13);path.lineTo(W-13,H-13);path.lineTo(W-13,H-104);path.lineTo(382,H-104);path.close();c.drawPath(path,fill=1,stroke=0)
  c.setStrokeColor(GOLD);c.setLineWidth(5);c.line(424,H-13,382,H-104)
  names=data['name'].split();first=' '.join(names[:-1]) or names[0];last=names[-1] if len(names)>1 else ''
  # Fit long legal names rather than allowing the sidebar to obscure them.
  font=30
  while c.stringWidth(first+' '+last,'Times-Bold',font)>365 and font>19:font-=1
  c.setFont('Times-Bold',font);c.setFillColor(NAVY);c.drawString(28,H-49,first)
  c.setFillColor(GOLD);c.drawString(28+c.stringWidth(first+' ','Times-Bold',font),H-49,last)
  para(data['target'].upper(),30,H-62,348,ParagraphStyle('target',parent=body,fontName='OTUBold',fontSize=8,leading=10))
  c.setFont('OTURegular',7);c.setFillColor(NAVY);c.drawString(30,H-96,'PEOPLE  ·  SYSTEMS  ·  OPPORTUNITY')
  para('Clear communication. Practical workflows. Organized follow-through.',430,H-35,145,side)
  c.setFillColor(HexColor('#F0F2F3'));c.roundRect(16,H-141,W-32,32,6,fill=1,stroke=0)
  para(data['contact'],29,H-118,W-58,ParagraphStyle('contact',parent=body,fontSize=8,leading=10))
  c.setFont('OTURegular',7);c.setFillColor(NAVY);c.drawString(28,22,data['name']);c.drawRightString(W-28,22,str(page))
  y=H-153
  if page==1:
   top=H-153
   tools=[x.strip() for x in re.split(r'[,;\n]',data.get('source_tools','')) if x.strip() and len(x.strip())<=100][:6]
   if tools:top=sidebar('TOOLS & TECHNOLOGIES',tools,top)
   top=sidebar('KEY CAPABILITIES',data['skills'][:4],top)
   if data['education']:top=sidebar('EDUCATION',[data['education']],top)

 def block(text,bold=False):
  nonlocal y
  style=ParagraphStyle('bold',parent=body,fontName='OTUBold') if bold else body
  p=Paragraph(html.escape(str(text)),style);_,h=p.wrap(367,700)
  if y-h<42:newpage()
  y=para(text,28,y,367,style)
 def title(text):
  nonlocal y
  if y<100:newpage()
  y=heading(text,28,y,367)
 newpage();title('PROFESSIONAL SUMMARY');block(data['summary']);title('CORE SKILLS & CAPABILITIES')
 # Two-column capability list follows the reference while preserving the actual wording.
 for i in range(0,len(data['skills']),2):
  heights=[]
  for col,text in enumerate(data['skills'][i:i+2]):
   x=28+col*188;c.setFillColor(PALE);c.circle(x+7,y-7,7,fill=1,stroke=0);c.setFillColor(GOLD);c.setFont('OTUBold',8);c.drawCentredString(x+7,y-10,'•')
   heights.append(para(text,x+20,y,163))
  y=min(heights)-2
 title('PROFESSIONAL EXPERIENCE')
 for role in data['roles']:
  if y<115:newpage();title('EXPERIENCE — CONTINUED')
  block(role['heading'],True)
  if role['dates']:block(role['dates'])
  for text in role['bullets'][:2]:block('• '+text)
  y-=5
 if data['projects']:
  title('FEATURED PROJECTS')
  for text in data['projects']:block(text)
 if data['recent_updates']:
  title('RECENT ACCOMPLISHMENTS')
  for text in data['recent_updates'][:3]:block(text)
 c.save();return out.getvalue()

def word(data,sections):
 from docx import Document
 from docx.shared import Inches,Pt,RGBColor
 from docx.oxml import OxmlElement
 from docx.oxml.ns import qn
 doc=Document();sec=doc.sections[0];sec.top_margin=sec.bottom_margin=Inches(.4);sec.left_margin=sec.right_margin=Inches(.35)
 normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(10);normal.paragraph_format.space_after=Pt(5)
 p=doc.add_paragraph();r=p.add_run(data['name']);r.font.name='Georgia';r.font.size=Pt(29);r.bold=True;r.font.color.rgb=RGBColor.from_string('09263B')
 p=doc.add_paragraph(data['target'].upper());p.runs[0].bold=True;p.runs[0].font.color.rgb=RGBColor.from_string('B89243')
 doc.add_paragraph(data['contact'])
 table=doc.add_table(rows=1,cols=2);table.autofit=False;table.columns[0].width=Inches(4.8);table.columns[1].width=Inches(2.9)
 left,right=table.rows[0].cells
 tcPr=right._tc.get_or_add_tcPr();shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'09263B');tcPr.append(shade)
 def add(cell,text,bold=False,gold=False):
  p=cell.add_paragraph();r=p.add_run(str(text));r.bold=bold;r.font.color.rgb=RGBColor.from_string('B89243' if gold else 'FFFFFF' if cell is right else '09263B');p.paragraph_format.keep_with_next=bold;return p
 for title,items in sections:
  if title in ('CORE CAPABILITIES','EDUCATION'):continue
  add(left,title,True,True)
  for item in items:
   if isinstance(item,dict):
    add(left,item['heading'],True)
    if item['dates']:add(left,item['dates'])
    for bullet in item['bullets'][:2]:add(left,'• '+bullet)
   else:add(left,item)
 add(right,'CAPABILITIES',True,True)
 for skill in data['skills']:add(right,'• '+skill)
 add(right,'EDUCATION',True,True);add(right,data['education'])
 if data.get('source_tools'):add(right,'TOOLS & SYSTEMS',True,True);add(right,data['source_tools'])
 out=io.BytesIO();doc.save(out);return out.getvalue()
